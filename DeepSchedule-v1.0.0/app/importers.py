# -*- coding: utf-8 -*-
"""把各种课表文件读成 timetable.json 用的结构。

支持三种来源：
1. 本项目的 `王の作息.xlsx`（每周一张 sheet，第 2 列是上课时间）
2. 教务系统导出的网页版课表（.xls / .html，其实是 <table>）
3. 自己填的 CSV（见 课表模板.csv）
"""
import csv
import io
import json
import re
from datetime import datetime

CN_NUM = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
          "十一", "十二", "十三", "十四"]

DEFAULT_PERIODS = [
    (1, "08:30", "09:15"), (2, "09:20", "10:05"), (3, "10:20", "11:05"),
    (4, "11:10", "11:55"), (5, "12:10", "12:55"), (6, "13:00", "13:45"),
    (7, "13:50", "14:35"), (8, "14:45", "15:30"), (9, "15:35", "16:20"),
    (10, "16:25", "17:10"), (11, "17:20", "18:05"), (12, "18:10", "18:55"),
    (13, "19:05", "19:50"), (14, "19:55", "20:40"),
]
PERIOD_TIME = {p: (s, e) for p, s, e in DEFAULT_PERIODS}

ENTRY_RE = re.compile(
    r"([^()]+?)\s*\(([^()]+)\)\s*\(([^()]+)\)\s*\(([^()]*(?:\([^()]*\))?[^()]*)\)")

DAY_CN = {"一": 0, "二": 1, "三": 2, "四": 3, "五": 4, "六": 5, "日": 6, "天": 6}
DAY_EN = {"mon": 0, "tue": 1, "wed": 2, "thu": 3, "fri": 4, "sat": 5, "sun": 6}


# ---------------------------------------------------------------- 工具
def parse_weeks(text):
    """'5-19' / '6-18双' / '5-19单' / '4-15,17-18' / '4,9' -> set[int]"""
    s = str(text).strip().replace(" ", "")
    parity = None
    if s.endswith("单"):
        parity, s = "odd", s[:-1]
    elif s.endswith("双"):
        parity, s = "even", s[:-1]
    weeks = set()
    for part in s.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-")[:2]
            seq = range(int(a), int(b) + 1)
        else:
            seq = [int(part)]
        for w in seq:
            if parity == "odd" and w % 2 == 0:
                continue
            if parity == "even" and w % 2 == 1:
                continue
            weeks.add(w)
    return weeks


def _norm_day(value):
    """'星期一' / '周一' / '1' / 'Mon' -> 0..6"""
    s = str(value).strip()
    if not s:
        return None
    if s.isdigit():
        n = int(s)
        return n - 1 if 1 <= n <= 7 else None
    m = re.search(r"(?:星期|周|礼拜)\s*([一二三四五六日天])", s)
    if m:
        return DAY_CN[m.group(1)]
    m = re.search(r"(\d)", s)
    if m:
        n = int(m.group(1))
        return n - 1 if 1 <= n <= 7 else None
    low = s.lower()[:3]
    return DAY_EN.get(low)


def _norm_period(value):
    s = str(value).strip()
    if s.isdigit():
        n = int(s)
        return n if 1 <= n <= 20 else None
    m = re.search(r"第\s*([一二三四五六七八九十]+)\s*节", s)
    if m and m.group(1) in CN_NUM:
        return CN_NUM.index(m.group(1)) + 1
    return None


def _add(weeks, week, day, period, name, room="", teacher="", times=None):
    if period is None or not name:
        return
    start, end = (times or PERIOD_TIME).get(period, ("", ""))
    row = {"p": period, "start": start, "end": end, "name": name.strip(),
           "room": (room or "").strip(), "teacher": (teacher or "").strip()}
    bucket = weeks.setdefault(str(week), {}).setdefault(str(day), [])
    bucket.append(row)


def _finalise(weeks):
    for days in weeks.values():
        for rows in days.values():
            rows.sort(key=lambda r: (r["p"], r["name"]))
    return weeks


# ---------------------------------------------------------------- 1. 课表工作簿
# 标准格式（本项目的 模版.xls / 王の作息.xlsx 都是这个）：
#   第 1 行表头：节次 | 上课时间 | 星期一 … 星期日
#   第 2 行起：第一节课 … （一行一节）
#   每个 sheet = 一周，sheet 名写成「第3周」或「3」都可以。
HEADER_WORDS = ("节次", "上课时间", "星期一", "周二", "星期")


def _sheet_week(name):
    """从 sheet 名里认周次：第3周 / 3 / Week 3 / 第三周。"""
    m = re.search(r"第\s*(\d+)\s*周", name)
    if m:
        return int(m.group(1))
    m = re.search(r"(?:week|wk)\s*(\d+)", name, re.I)
    if m:
        return int(m.group(1))
    if name.strip().isdigit():
        return int(name.strip())
    m = re.search(r"第\s*([一二三四五六七八九十]+)\s*周", name)
    if m and m.group(1) in CN_NUM:
        return CN_NUM.index(m.group(1)) + 1
    return None


def _find_header(ws, limit=6):
    """找表头行，返回 (行号, {列号: 'period'|'time'|'day0'…})。"""
    for row in range(1, min(ws.max_row, limit) + 1):
        cells = [(col, str(ws.cell(row=row, column=col).value or "").strip())
                 for col in range(1, min(ws.max_column, 12) + 1)]
        joined = " ".join(v for _, v in cells)
        if not any(word in joined for word in HEADER_WORDS):
            continue
        mapping = {}
        for col, text in cells:
            if not text:
                continue
            if "节次" in text:
                mapping[col] = "period"
            elif "时间" in text:
                mapping[col] = "time"
            else:
                day = _norm_day(text)
                if day is not None:
                    mapping[col] = f"day{day}"
        if "period" in mapping.values() and any(v.startswith("day") for v in mapping.values()):
            return row, mapping
    return None, {}


def from_workbook(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    weeks = {}
    for sheet in wb.sheetnames:
        week = _sheet_week(sheet)
        if week is None:
            continue
        ws = wb[sheet]
        header_row, mapping = _find_header(ws)
        if header_row is None:
            continue
        period_col = next((c for c, k in mapping.items() if k == "period"), 1)
        time_col = next((c for c, k in mapping.items() if k == "time"), None)
        day_cols = {int(k[3:]): c for c, k in mapping.items() if k.startswith("day")}

        for row_idx in range(header_row + 1, min(ws.max_row, header_row + 24) + 1):
            period = _norm_period(ws.cell(row=row_idx, column=period_col).value)
            if period is None:
                continue
            times = None
            if time_col:
                tm = re.match(r"(\d{1,2}:\d{2})\s*[-~－—]\s*(\d{1,2}:\d{2})",
                              str(ws.cell(row=row_idx, column=time_col).value or ""))
                if tm:
                    times = {period: (tm.group(1), tm.group(2))}
                    PERIOD_TIME.setdefault(period, (tm.group(1), tm.group(2)))
            for day, col in day_cols.items():
                text = _cell_text(ws.cell(row=row_idx, column=col).value)
                if not text:
                    continue
                parts = [p.strip() for p in text.replace("\r", "").split("\n") if p.strip()]
                name = parts[0] if parts else ""
                room = parts[1] if len(parts) > 1 else ""
                teacher = parts[2] if len(parts) > 2 else ""
                _add(weeks, week, day, period, name, room, teacher, times)
    if not weeks:
        raise ValueError("没读懂这个工作簿：需要「节次/上课时间/星期一…」表头，且 sheet 名带周次")
    return {"weeks": _finalise(weeks), "meta": {"source": "xlsx"}}


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    try:                                    # openpyxl 的富文本
        return "".join(getattr(part, "text", str(part)) for part in value).strip()
    except TypeError:
        return str(value).strip()


# ---------------------------------------------------------------- 2. 教务系统网页课表
def _decode(raw):
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "ignore")


def from_html(path):
    import pandas as pd

    with open(path, "rb") as fh:
        html = _decode(fh.read())
    tables = pd.read_html(io.StringIO(html))
    target = None
    for table in tables:
        labels = [str(v) for v in table.iloc[:, 0].tolist()]
        if sum(1 for v in labels if re.search(r"第.+节", v)) >= 3:
            target = table
            break
    if target is None:
        raise ValueError("没找到「节次 × 星期」的课表表格")

    # read_html 已经把表头那一行吃成列名了：列名是「星期一…星期日」
    day_cols = {}
    for pos, label in enumerate(target.columns):
        day = _norm_day(label)
        if day is not None and pos > 0:
            day_cols[pos] = day
    if len(day_cols) < 5:                    # 表头没认出来就按 1..7 顺排
        day_cols = {pos: pos - 1 for pos in range(1, min(8, target.shape[1]))}

    weeks = {}
    for row in range(target.shape[0]):
        label = str(target.iat[row, 0])
        period = _norm_period(label)
        if period is None:
            continue
        for col, day in day_cols.items():
            value = target.iat[row, col]
            if not isinstance(value, str) or not value.strip():
                continue
            for m in ENTRY_RE.finditer(value):
                name, _code, teacher, tail = (m.group(1).strip(), m.group(2).strip(),
                                              m.group(3).strip(), m.group(4).strip())
                tm = re.match(r"(\S+)\s+(.*)$", tail)
                if not tm:
                    continue
                for week in parse_weeks(tm.group(1)):
                    _add(weeks, week, day, period, name, tm.group(2), teacher)
    if not weeks:
        raise ValueError("表格认出来了，但没解析出任何课程")
    return {"weeks": _finalise(weeks), "meta": {"source": "html"}}


# ---------------------------------------------------------------- 2.5 SpreadsheetML
def _header_map(cells):
    """['节次','上课时间','星期一'…] -> {列号: 'period'|'time'|'day0'…}"""
    mapping = {}
    for col, text in enumerate(cells):
        text = (text or "").strip()
        if not text:
            continue
        if "节次" in text:
            mapping[col] = "period"
        elif "时间" in text:
            mapping[col] = "time"
        else:
            day = _norm_day(text)
            if day is not None:
                mapping[col] = f"day{day}"
    return mapping


def _consume_table(weeks, week, rows):
    """rows = [[单元格文本…], …]，第一行是表头。"""
    mapping = {}
    header_row = 0
    for i, row in enumerate(rows[:4]):
        m = _header_map(row)
        if "period" in m.values() and any(v.startswith("day") for v in m.values()):
            mapping, header_row = m, i
            break
    if not mapping:
        return False
    period_col = next(c for c, k in mapping.items() if k == "period")
    time_col = next((c for c, k in mapping.items() if k == "time"), None)
    day_cols = {int(k[3:]): c for c, k in mapping.items() if k.startswith("day")}
    for row in rows[header_row + 1:]:
        def at(idx):
            return row[idx] if idx is not None and idx < len(row) else ""
        period = _norm_period(at(period_col))
        if period is None:
            continue
        times = None
        if time_col is not None:
            tm = re.match(r"(\d{1,2}:\d{2})\s*[-~－—]\s*(\d{1,2}:\d{2})", at(time_col))
            if tm:
                times = {period: (tm.group(1), tm.group(2))}
                PERIOD_TIME.setdefault(period, (tm.group(1), tm.group(2)))
        for day, col in day_cols.items():
            text = (at(col) or "").replace("\r", "").strip()
            if not text:
                continue
            parts = [x.strip() for x in text.split("\n") if x.strip()]
            name = parts[0] if parts else ""
            room = parts[1] if len(parts) > 1 else ""
            teacher = parts[2] if len(parts) > 2 else ""
            _add(weeks, week, day, period, name, room, teacher, times)
    return True


def from_spreadsheetml(path):
    """Excel 2003 XML 表格（模版.xls 就是这种）：UTF-8，可以有多张 worksheet。"""
    import xml.etree.ElementTree as ET

    ns = "{urn:schemas-microsoft-com:office:spreadsheet}"
    root = ET.parse(path).getroot()
    weeks = {}
    for sheet in root.findall(f"{ns}Worksheet"):
        week = _sheet_week(sheet.get(f"{ns}Name") or "")
        if week is None:
            continue
        table = sheet.find(f"{ns}Table")
        if table is None:
            continue
        rows = []
        for row in table.findall(f"{ns}Row"):
            cells = []
            for cell in row.findall(f"{ns}Cell"):
                index = cell.get(f"{ns}Index")
                if index:
                    while len(cells) < int(index) - 1:
                        cells.append("")
                data = cell.find(f"{ns}Data")
                cells.append((data.text or "") if data is not None else "")
            rows.append(cells)
        _consume_table(weeks, week, rows)
    if not weeks:
        raise ValueError("这个 XML 表格里没有可用的周次工作表")
    return {"weeks": _finalise(weeks), "meta": {"source": "spreadsheetml"}}


# ---------------------------------------------------------------- 3. CSV
CSV_COLUMNS = ["星期", "节次", "课程", "教室", "教师", "周次"]


def from_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise ValueError("CSV 是空的")
    head = set(rows[0].keys())
    missing = [c for c in ("星期", "节次", "课程", "周次") if c not in head]
    if missing:
        raise ValueError("CSV 缺少列：" + "、".join(missing))

    weeks = {}
    for row in rows:
        day = _norm_day(row.get("星期"))
        period = _norm_period(row.get("节次"))
        name = (row.get("课程") or "").strip()
        if day is None or period is None or not name:
            continue
        for week in parse_weeks(row.get("周次") or ""):
            _add(weeks, week, day, period, name,
                 row.get("教室") or "", row.get("教师") or "")
    if not weeks:
        raise ValueError("CSV 里没有可用的行（检查「星期/节次/周次」）")
    return {"weeks": _finalise(weeks), "meta": {"source": "csv"}}


# ---------------------------------------------------------------- 入口
def load_any(path):
    lower = path.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return from_workbook(path)
    if lower.endswith(".csv"):
        return from_csv(path)
    with open(path, "rb") as fh:
        head = fh.read(4096).lstrip()
    if b"urn:schemas-microsoft-com:office:spreadsheet" in head:
        return from_spreadsheetml(path)
    if head.startswith(b"<") or b"<table" in head.lower() or lower.endswith((".html", ".htm")):
        return from_html(path)
    if lower.endswith(".xls"):
        return from_html(path)      # 教务系统的 .xls 其实就是网页
    return from_html(path)


if __name__ == "__main__":
    import sys
    data = load_any(sys.argv[1])
    weeks = sorted(int(w) for w in data["weeks"])
    total = sum(len(r) for d in data["weeks"].values() for r in d.values())
    print("weeks", weeks[:3], "...", weeks[-1], "count", len(weeks), "entries", total)
